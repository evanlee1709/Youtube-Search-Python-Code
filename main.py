import os
import logging
import scrapetube
from googleapiclient.discovery import build
from googleapiclient.errors import HttpError
from openpyxl import Workbook, load_workbook
from langchain_community.llms import Ollama
from youtube_transcript_api import YouTubeTranscriptApi as yta

API_KEY = 'AIzaSyBk-J3xo7kVE8TXxWfSkI2PLHhB_Du_oNA'
youtube = build('youtube', 'v3', developerKey=API_KEY)

llm = Ollama(base_url="http://131.123.41.132:11434", model="llava:34b")

# Set up logging
logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')


def get_video_details(video_id):
    """Fetch video details such as title, description, and publish date using the YouTube Data API."""
    try:
        request = youtube.videos().list(part="snippet, statistics", id=video_id)
        response = request.execute()
        items = response.get('items', [])
        if items:
            snippet = items[0].get('snippet', {})
            statistics = items[0].get('statistics', {})
            title = snippet.get('title', '')
            description = snippet.get('description', '')
            publish_date = snippet.get('publishedAt', '')
            views = statistics.get('viewCount', '')
            return title, description, publish_date, views
    except HttpError as e:
        logging.error(f"API error for video {video_id}: {e}")
    except Exception as e:
        logging.error(f"Error fetching details for video {video_id}: {e}")
    return '', '', '', ''


def get_transcript(video_id):
    try:
        data = yta.get_transcript(video_id)
        transcript = ' '.join([val['text'] for val in data if 'text' in val])
        return transcript
    except Exception as e:
        logging.error(f"Error fetching transcript for video {video_id}: {e}")
        return ''


def organize_transcript(transcript):
    if not transcript:
        return "Transcript does not exist"

    summary_text = llm.invoke(f"""
                              edit the following video transcript with proper punctuation and capitalization:
                              "{transcript}"
                              """)
    return summary_text

def summarize_transcript(transcript):
    if not transcript:
        return "Transcript does not exist"

    summary_text = llm.invoke(f"""
                              please give me a summary of the following text:
                              "{transcript}"
                              """)
    return summary_text

def save_to_excel(data, filename='AI DataLog_Transcript_Record.xlsx'):
    """Save data to an Excel file."""
    try:
        if os.path.exists(filename):
            wb = load_workbook(filename)
            ws = wb.active
        else:
            wb = Workbook()
            ws = wb.active
            ws.append(['Title', 'Video ID', 'URL', 'Description', 'Publish Date', 'Views', 'Edited Transcript',
                       'Transcript Summary'])

        for row in data:
            ws.append(row)

        wb.save(filename)
    except Exception as e:
        logging.error(f"Error saving data to Excel: {e}")


def track_progress(current, total):
    logging.info(f"Processed {current}/{total} videos. {total - current} videos left.")


def remove_duplicates(videos):
    seen = set()
    unique_videos = []
    for video in videos:
        video_id = video['videoId']
        if video_id not in seen:
            seen.add(video_id)
            unique_videos.append(video)
    return unique_videos


def main():
    search_query = input("YouTube Search: ")
    total_videos = int(input("How many videos do you want to analyze? "))
    include_no_transcript = input("Do you want to include videos without transcripts? (yes/no): ").strip().lower() == 'yes'
    all_data = []

    videos_generator = scrapetube.get_search(search_query, limit=total_videos)
    videos = list(videos_generator)
    unique_videos = remove_duplicates(videos)

    for i, video in enumerate(unique_videos):
        video_id = video['videoId']
        video_url = f'https://www.youtube.com/watch?v={video_id}'
        title, description, publish_date, views = get_video_details(video_id)

        transcript = get_transcript(video_id)
        if not transcript and not include_no_transcript:
            continue

        edited_transcript = organize_transcript(transcript) if transcript else "Transcript not available"
        summary = summarize_transcript(transcript) if transcript else "Summary not available"

        # Print details in the specified order
        print("\nURL:", video_url)
        print("Title:", title)
        print("Publish Date:", publish_date)
        print("Views:", views)
        print("Description:", description)
        print("Transcript:", edited_transcript)
        print("Transcript Summary:", summary)

        data = [title, video_id, video_url, description, publish_date, views, edited_transcript, summary]
        all_data.append(data)

        # Track progress
        track_progress(i + 1, len(unique_videos))

        if i >= total_videos - 1:
            break

    save_to_excel(all_data)


if __name__ == "__main__":
    main()
